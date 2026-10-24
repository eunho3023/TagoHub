"""서울교통공사 공공데이터(역 연락처·승강설비·경사로·수유실·화장실·물품보관함·환승거리 …)를 모아
앱의 역 '편의정보' 자료(app/src/main/assets/subway_station_facilities.tsv)에 서울교통공사 행으로 보탠다.

  python pipeline/build_seoul_metro_info.py [파일이 있는 폴더]     # 기본: ~/Downloads

여러 번 돌려도 같은 결과가 되도록, 이 스크립트가 만드는 항목(CATEGORIES)의 서울교통공사 행은 먼저 지우고 다시 씀.
파일 이름은 앞부분으로 찾음. 없는 파일은 건너뜀. 한 줄 = 운영사 \t 노선 \t 역 \t 항목 \t 내용 (내용 안의 구분은 ' · ').
"""
import csv
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "app/src/main/assets/subway_station_facilities.tsv"
OPERATOR = "서울교통공사"

CATEGORIES = [
    "역 연락처", "환승 거리·시간", "환승 안내", "승강설비", "휠체어 경사로", "수유실", "무인민원발급기", "물품보관함",
    "전동휠체어 충전기", "자전거 보관소", "화장실", "시각장애인 음성유도기", "무빙워크", "편의시설", "약냉방칸",
]


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("cp949", "replace")


def find(folder: Path, start: str):
    for p in sorted(folder.iterdir()):
        if p.name.startswith(start):
            return p
    print("  (없음)", start)
    return None


def csv_rows(folder, start):
    p = find(folder, start)
    if not p:
        return []
    rows = list(csv.reader(read_text(p).splitlines()))
    head = [h.strip() for h in rows[0]]
    return [{head[i]: (r[i].strip() if i < len(r) else "") for i in range(len(head))} for r in rows[1:] if any(c.strip() for c in r)]


def xlsx_rows(folder, start, sheet, header_contains):
    p = find(folder, start)
    if not p:
        return []
    ws = openpyxl.load_workbook(p, read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    def norm(c):
        return re.sub(r"\s", "", str(c)) if c is not None else ""   # '역  명' 처럼 가운데 공백·줄바꿈이 있는 머리글도 맞춤

    hi = next(i for i, r in enumerate(rows) if r and header_contains in [norm(c) for c in r])
    head = [norm(c) for c in rows[hi]]
    return [{head[i]: (str(r[i]).strip() if r[i] is not None else "") for i in range(len(head))} for r in rows[hi + 1:] if r and any(c is not None for c in r)]


def line_name(v: str):
    """'1' / '01호선' / '1호선' → '1호선' (1~9 만, 아니면 None)"""
    m = re.match(r"\s*0*([1-9])(\.0)?\s*(호선)?\s*$", str(v))   # 엑셀 숫자 칸은 '3.0' 으로 읽힘
    return f"{m.group(1)}호선" if m else None


def clean(name: str) -> str:
    """'서울역 (1)', '종로3가역(1호선)' → '서울역', '종로3가역' (앱이 역 이름을 맞출 때 괄호는 떼므로)"""
    return re.sub(r"\s*\(.*?\)", "", str(name)).strip()


facts = defaultdict(list)  # (line, station, category) -> [detail]


def add(line, station, category, detail):
    ln = line_name(line)
    st = clean(station)
    detail = re.sub(r"\s+", " ", str(detail)).strip(" ·")
    if ln and st and detail:
        if detail not in facts[(ln, st, category)]:
            facts[(ln, st, category)].append(detail)


def mmss(v: str) -> str:
    m = re.match(r"(\d+):(\d+)", v or "")
    if not m:
        return v or ""
    mi, se = int(m.group(1)), int(m.group(2))
    return f"{mi}분" + (f" {se}초" if se else "")


def main():
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Downloads"

    stations = defaultdict(set)  # 호선 → 역들 (약냉방칸을 노선의 모든 역에 붙이려고)
    for r in csv_rows(folder, "서울교통공사_역주소 및 전화번호"):
        add(r["호선"], r["역명"], "역 연락처", f"전화 {r['역전화번호']} · 도로명주소: {r['도로명주소']}")
        ln = line_name(r["호선"])
        if ln:
            stations[ln].add(clean(r["역명"]))

    for r in csv_rows(folder, "서울교통공사_환승역거리 소요시간"):
        add(r["호선"], r["환승역명"], "환승 거리·시간", f"{r['환승노선']} 환승 · {r['환승거리']}m · {mmss(r['환승소요시간'])}")

    for r in csv_rows(folder, "서울교통공사_수도권 도시철도 환승 데이터"):
        add(r["환승시작 호선"], r["환승시작역"], "환승 안내",
            f"{r['환승종료 호선']} 환승 ({r['환승 열차 방면']}) · 내리는 곳 {r['하차위치(호차)']}호차 {r['하차위치(문)']}번 문 ({r['하차 열차 방면']}) → "
            f"타는 곳 {r['환승 승차위치(호차)']}호차 {r['환승 승차위치(문)']}번 문 · {mmss(r['소요시간'])}")

    line = ""
    for r in xlsx_rows(folder, "노약자장애인_편의시설", "역별승강설비현황", "역명"):
        line = r.get("호선") or line
        if not r.get("역명"):
            continue
        parts = [f"{label} {int(float(r[k]))}대" for label, k in
                 (("엘리베이터", "E/L"), ("에스컬레이터", "E/S"), ("무빙워크", "M/W"), ("휠체어리프트", "W/L"), ("수직형리프트", "수직형리프트"))
                 if r.get(k) and r[k] not in ("0", "0.0")]
        add(line, r["역명"], "승강설비", " · ".join(parts))

    ramps = defaultdict(list)
    for r in csv_rows(folder, "서울교통공사 휠체어경사로"):
        ramps[(r["호선"], r["역명"])].append(f"{r['구분']} {r['위치']}".strip())
    for (ln, st), items in ramps.items():
        add(ln, st, "휠체어 경사로", "설치 위치: " + ", ".join(items))

    for r in csv_rows(folder, "서울교통공사_수유실현황"):
        add(r["호선"], r["역명"], "수유실", f"{r['상세위치']} · {r['시설구분']} · {r['면적(제곱미터)']}㎡ · {r['운임 비운임']}")

    for r in csv_rows(folder, "서울교통공사_무인민원발급기현황"):
        add(r["호선"], r["역명"], "무인민원발급기", f"{r['상세위치']} ({r['지상지하구분']} {r['역층']}) · {r['설치대수']}대 · {r['계약자(업체명)']} {r['전화번호']}")

    for r in csv_rows(folder, "서울교통공사_물품보관함"):
        st = re.sub(r"[\d~\-]+$", "", r["보관함명"]).strip()
        add(r["호선"], st, "물품보관함", f"{r['보관함명']} · {r['상세위치']}")

    for r in csv_rows(folder, "서울교통공사_전동휠체어 급속충전기"):
        add(r["호선"], r["역명"], "전동휠체어 충전기", f"{r['설치위치']} · {r['충전기수']}대")

    bikes = defaultdict(list)
    for r in csv_rows(folder, "서울교통공사_1_8호선 자전거 보관소"):
        bikes[(r["호선"], r["역명"])].append(f"{r['위치정보(출입구번호)']}번 출입구 {r['대수']}대")
    for (ln, st), items in bikes.items():
        add(ln, st, "자전거 보관소", ", ".join(items))

    for start in ("서울교통공사_역사공중화장실정보", "서울교통공사_역사장애인화장실정보"):
        for r in csv_rows(folder, start):
            add(r["운영노선명"], r["역명"], "화장실", f"{r['구분']} · {r['상세위치']} ({r['지상 또는 지하 구분']} {r['역층']}층) · {r['개방시간']}")
    for start in ("서울교통공사_9호선2_3단계 역사공중화장실정보", "서울교통공사_9호선2_3단계 역사장애인화장실정보"):
        for r in csv_rows(folder, start):
            add(r["호선"], r["화장실명(역명)"], "화장실", f"{r['구분']} · {r['화장실 상세위치']} ({r['지상-지하 구분'] if '지상-지하 구분' in r else r.get('지상_지하 구분', '')} {r['역층']}층) · {r['개방시간']}")

    guides = defaultdict(list)
    for r in csv_rows(folder, "서울교통공사 지하철 시각장애인 음성유도기"):
        guides[(r["호선"], r["역명"])].append(r["설치위치"])
    for (ln, st), items in guides.items():
        sample = ", ".join(dict.fromkeys(items[:4]))
        add(ln, st, "시각장애인 음성유도기", f"{len(items)}곳 설치 · 예: {sample}")

    ln_, st_ = "", ""
    for r in xlsx_rows(folder, "서울교통공사_무빙워크 설치현황", "무빙워크 설치현황", "역명"):
        ln_, st_ = r.get("호선") or ln_, r.get("역명") or st_      # 합쳐진 칸은 첫 줄에만 호선·역명이 있음
        add(ln_, st_, "무빙워크", f"{r['설치위치']} · {r['호기']}호기")

    flags = [("문화공간여부", "문화공간"), ("휠체어리프트여부", "휠체어리프트"), ("만남의장소여부", "만남의 장소"), ("환승주차장여부", "환승주차장"),
             ("자전거보관소여부", "자전거보관소"), ("엘리베이터여부", "엘리베이터"), ("무인민원발급기여부", "무인민원발급기"),
             ("환전키오스크여부", "환전 키오스크"), ("수유실여부", "수유실")]
    for r in csv_rows(folder, "서울교통공사_편의시설 현황"):
        have = [label for key, label in flags if r.get(key) == "Y"]
        if have:
            add(r["호선"], r["역명"], "편의시설", ", ".join(have))

    p = find(folder, "서울교통공사_호선별 약냉방칸")
    if p:
        ws = openpyxl.load_workbook(p, read_only=True, data_only=True).worksheets[0]
        cars = defaultdict(list)
        for row in list(ws.iter_rows(values_only=True))[1:]:
            if row and row[1] and row[2]:
                cars[line_name(row[1])].append(str(row[2]))
        for ln, lst in cars.items():
            for st in stations.get(ln, ()):
                add(ln, st, "약냉방칸", ", ".join(lst) + " (호선 공통)")

    # 기존 파일에서 이 스크립트의 서울교통공사 행을 지우고 새로 씀
    lines = OUT.read_text(encoding="utf-8-sig").splitlines()
    header, body = lines[0], lines[1:]
    keep = [l for l in body if not (l.split("\t")[0] == OPERATOR and len(l.split("\t")) > 3 and l.split("\t")[3] in CATEGORIES)]
    new = [f"{OPERATOR}\t{ln}\t{st}\t{cat}\t{d}" for (ln, st, cat), ds in sorted(facts.items()) for d in ds]
    OUT.write_text("﻿" + "\n".join([header] + keep + new) + "\n", encoding="utf-8")
    by_cat = defaultdict(int)
    for (_, _, cat), ds in facts.items():
        by_cat[cat] += len(ds)
    print(f"기존 {len(keep)}행 유지, 새로 {len(new)}행")
    for c in CATEGORIES:
        print(f"  {c}: {by_cat[c]}")


if __name__ == "__main__":
    main()
